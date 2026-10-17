"""
smg_lab.py -- experiment engine used by SMG_Lite_Experiments.ipynb

Every notebook cell calls ONE function from here, e.g.

    run("plantvillage", "MobileNetV3-Large")        # one model on one dataset
    run("rice", "SMG_Lite", aug=False)               # no-augmentation run (training-curve figure)
    run_ablation("rice", ghost=True, msdw=False, cea=True)

Each call trains on the TRAIN split, picks the best epoch on the VALIDATION
split and reports all metrics on the held-out TEST split.  Results are saved to
disk, so a finished cell is skipped when re-run, and an interrupted cell resumes
from its last epoch.  The "combine" functions at the bottom read those saved
results and produce the tables / figures / filled Word file in the layout of
templates/result_word.docx.

The model file SMG_Lite.py is imported unchanged.  The training recipe re-uses
train.py (MixUp/CutMix + EMA) with the same augmentation and optimiser settings.
"""

import os, json, math, time, copy, random, warnings

import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.utils.data import Dataset, DataLoader
from torchvision import models, transforms
from torchvision.transforms import RandAugment
from PIL import Image

from SMG_Lite import smg_lite, count_parameters, GhostExpand, CoordECA
from benchmark import macs_of
from train import mix_batch, EMA

# --------------------------------------------------------------------------- #
#  configuration (overwritten by configure() from the notebook)
# --------------------------------------------------------------------------- #
CFG = dict(
    datasets={},                 # name -> dict(path=..., title=..., epochs=...)
    out_dir="outputs",
    split=(0.70, 0.15, 0.15),    # train / validation / test
    seed=42,
    img_size=224, batch_size=64, workers=4,
    lr=3e-3, wd=5e-2, warmup=3,
    mixup=0.2, smooth=0.1, ema_decay=0.999,
    randaug_magnitude=7,
    mixup_off_frac=0.2,          # MixUp/CutMix switched off for the last 20 % of epochs
    tta=True,                    # horizontal-flip test-time augmentation (applied to EVERY model)
    lr_override={"VGG16": 1e-4}, # plain VGG16 (no BatchNorm) diverges at 3e-3 from scratch
    transfer_from=None,          # e.g. "plantvillage": init rice/mango runs from that dataset's checkpoint
    ablation_epochs=20,
    force=False,                 # True -> retrain even if results exist
    max_batches=None,            # int -> limit batches per epoch (quick smoke test only)
    backup_dir=None,             # e.g. Google Drive folder: every result/checkpoint is mirrored there
)

MODEL_NAMES = ["MobileNetV3-Large", "MobileNetV2", "VGG16", "ResNet50",
               "InceptionV3", "Xception", "DenseNet121", "SMG_Lite"]

# row labels exactly as in result_word.docx
DISPLAY = {"MobileNetV3-Large": "MobileNet V3 Large", "MobileNetV2": "MobileNet V2",
           "VGG16": "Standard VGG16", "ResNet50": "ResNet50", "InceptionV3": "InceptionV3",
           "Xception": "Xception", "DenseNet121": "DenseNet121", "SMG_Lite": "SMG_Lite"}

TABLE_COLS = ["Model", "Model Size (in MB)", "Total Trainable parameters", "Loss",
              "Accuracy", "Precision", "Recall", "F1-score", "AUC- Score",
              "Kappa Score", "GFLOPS"]

ABLATION_VARIANTS = [  # (ghost, ms-dw gate, cea) -- row order of the ablation table
    (False, False, False), (True, False, False), (False, True, False), (False, False, True),
    (True, True, False), (True, False, True), (False, True, True), (True, True, True)]

DEVICE = "cuda" if torch.cuda.is_available() else "cpu"


def configure(**kw):
    CFG.update(kw)
    os.makedirs(CFG["out_dir"], exist_ok=True)
    if DEVICE == "cuda":
        torch.backends.cudnn.benchmark = True
    if CFG.get("backup_dir"):
        restore_from_backup()
    return CFG


# --------------------------------------------------------------------------- #
#  FAST + RELIABLE: train on the local disk, mirror every result to a backup
#  folder (Google Drive) and restore it automatically in a new session
# --------------------------------------------------------------------------- #
def _backup_path(local):
    rel = os.path.relpath(local, CFG["out_dir"])
    return os.path.join(CFG["backup_dir"], rel)


def _backup(local):
    """Copy a file (or every file of a folder) from out_dir to backup_dir.
    Writes to a temporary name first, so an interrupted copy never damages the backup."""
    import shutil
    if not CFG.get("backup_dir") or not os.path.exists(local):
        return
    files = [local] if os.path.isfile(local) else \
        [os.path.join(r, f) for r, _, fs in os.walk(local) for f in fs]
    for f in files:
        dst = _backup_path(f)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        try:
            shutil.copyfile(f, dst + ".tmp")
            os.replace(dst + ".tmp", dst)
        except OSError as e:
            print(f"  (backup of {f} failed: {e}; training continues)")


def _unbackup(local):
    if CFG.get("backup_dir") and os.path.exists(_backup_path(local)):
        os.remove(_backup_path(local))


def restore_from_backup():
    """New session: copy everything saved in backup_dir back to out_dir (only missing/changed files)."""
    import shutil
    src = CFG["backup_dir"]
    if not os.path.isdir(src):
        os.makedirs(src, exist_ok=True)
        print(f"backup folder created: {src}")
        return
    n = 0
    for r, _, fs in os.walk(src):
        for f in fs:
            if f.endswith(".tmp"):
                continue
            s_ = os.path.join(r, f)
            d_ = os.path.join(CFG["out_dir"], os.path.relpath(s_, src))
            if not os.path.exists(d_) or os.path.getsize(d_) != os.path.getsize(s_):
                os.makedirs(os.path.dirname(d_), exist_ok=True)
                shutil.copyfile(s_, d_)
                n += 1
    print(f"restored {n} file(s) from backup {src} -> {CFG['out_dir']}")


def prepare_data(zip_dir=None):
    """Make sure every dataset is on the LOCAL disk (fast).  If data/<name> has no images and
    <zip_dir>/<name>.zip exists (e.g. on Google Drive), unzip it there once per session."""
    import zipfile
    for name, d in CFG["datasets"].items():
        path = d["path"]
        has_imgs = os.path.isdir(path) and any(
            os.path.splitext(f)[1].lower() in IMG_EXT for _, _, fs in os.walk(path) for f in fs)
        if has_imgs:
            print(f"{name:13s} ready  ({path})")
            continue
        z = os.path.join(zip_dir, f"{name}.zip") if zip_dir else None
        if z and os.path.exists(z):
            t0 = time.time()
            os.makedirs(path, exist_ok=True)
            with zipfile.ZipFile(z) as zf:
                zf.extractall(path)
            print(f"{name:13s} unzipped {z} -> {path}  ({time.time()-t0:.0f}s)")
        else:
            print(f"{name:13s} MISSING: put images in {path}"
                  + (f" or upload {z}" if z else ""))


def _p(*parts):
    path = os.path.join(CFG["out_dir"], *parts)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    return path


def seed_all(seed):
    random.seed(seed); np.random.seed(seed); torch.manual_seed(seed)
    torch.cuda.manual_seed_all(seed)


# --------------------------------------------------------------------------- #
#  data: stratified train / validation / test split (saved, so it never changes)
# --------------------------------------------------------------------------- #
IMG_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}


def _junk(name):
    return name.startswith((".", "__MACOSX", "._"))


def _subdirs(root):
    return [d for d in os.listdir(root) if os.path.isdir(os.path.join(root, d)) and not _junk(d)]


def _scan(folder):
    items = []
    for cls in sorted(_subdirs(folder)):
        for r, ds, fs in os.walk(os.path.join(folder, cls)):
            ds[:] = [x for x in ds if not _junk(x)]
            for f in sorted(fs):
                if not _junk(f) and os.path.splitext(f)[1].lower() in IMG_EXT:
                    items.append((os.path.join(r, f), cls))
    return items


def _resolve_root(root):
    while True:
        subs = {d.lower(): d for d in _subdirs(root)}
        if "color" in subs and ("grayscale" in subs or "segmented" in subs):  # raw PlantVillage
            return os.path.join(root, subs["color"])
        loose = [f for f in os.listdir(root) if os.path.splitext(f)[1].lower() in IMG_EXT]
        if len(subs) == 1 and not loose:           # zip with an extra top folder -> go inside
            root = os.path.join(root, next(iter(subs.values())))
            continue
        return root


def get_split(name):
    cache = _p("splits", f"{name}.json")
    if os.path.exists(cache):
        with open(cache) as f:
            out = json.load(f)
        # same images (relative paths), but read them from wherever the dataset is NOW
        # (e.g. unzipped to /content in a new Colab session)
        out["root"] = _resolve_root(CFG["datasets"][name]["path"])
        return out
    from sklearn.model_selection import train_test_split
    root = _resolve_root(CFG["datasets"][name]["path"])
    subs = {d.lower(): d for d in _subdirs(root)}
    vkey = next((k for k in ("val", "valid", "validation") if k in subs), None)

    if "train" in subs and vkey and "test" in subs:            # already split -> use it
        parts = {s: _scan(os.path.join(root, subs[k]))
                 for s, k in (("train", "train"), ("val", vkey), ("test", "test"))}
        classes = sorted({c for _, c in parts["train"]})
    else:                                                      # split it ourselves
        keys = [k for k in subs if k in ("train", "val", "valid", "validation", "test")]
        items = sum([_scan(os.path.join(root, subs[k])) for k in keys], []) if keys else _scan(root)
        classes = sorted({c for _, c in items})
        y = [c for _, c in items]
        tr, va, te = CFG["split"]
        rest, test = train_test_split(items, test_size=te, stratify=y, random_state=CFG["seed"])
        train, val = train_test_split(rest, test_size=va / (tr + va),
                                      stratify=[c for _, c in rest], random_state=CFG["seed"])
        parts = {"train": train, "val": val, "test": test}

    cidx = {c: i for i, c in enumerate(classes)}
    out = {"root": root, "classes": classes}
    for s, its in parts.items():
        out[s] = [[os.path.relpath(p, root), cidx[c]] for p, c in its]
    with open(cache, "w") as f:
        json.dump(out, f)
    _backup(cache)
    print(f"[{name}] {len(classes)} classes | train {len(out['train'])} | "
          f"val {len(out['val'])} | test {len(out['test'])}  -> split saved to {cache}")
    return out


class ImageList(Dataset):
    def __init__(self, root, items, tf):
        self.root, self.items, self.tf = root, items, tf

    def __len__(self):
        return len(self.items)

    def __getitem__(self, i):
        p, y = self.items[i]
        img = Image.open(os.path.join(self.root, p)).convert("RGB")
        return self.tf(img), y


def _transforms(aug):
    s = CFG["img_size"]
    mean, std = [0.485, 0.456, 0.406], [0.229, 0.224, 0.225]
    eval_tf = transforms.Compose([           # identical to train.py
        transforms.Resize(int(s * 1.15)), transforms.CenterCrop(s),
        transforms.ToTensor(), transforms.Normalize(mean, std)])
    if not aug:
        return eval_tf, eval_tf
    train_tf = transforms.Compose([          # identical to train.py
        transforms.RandomResizedCrop(s, scale=(0.6, 1.0)),
        transforms.RandomHorizontalFlip(), transforms.RandomVerticalFlip(),
        transforms.RandomRotation(25),
        transforms.ColorJitter(0.25, 0.25, 0.25, 0.05),
        RandAugment(num_ops=2, magnitude=CFG["randaug_magnitude"]),
        transforms.ToTensor(), transforms.Normalize(mean, std),
        transforms.RandomErasing(p=0.25)])
    return train_tf, eval_tf


def _loaders(split, aug):
    train_tf, eval_tf = _transforms(aug)
    kw = dict(num_workers=CFG["workers"], pin_memory=(DEVICE == "cuda"),
              persistent_workers=CFG["workers"] > 0)
    bs = CFG["batch_size"]
    tr = DataLoader(ImageList(split["root"], split["train"], train_tf), bs, shuffle=True,
                    drop_last=len(split["train"]) > bs, **kw)
    va = DataLoader(ImageList(split["root"], split["val"], eval_tf), bs, shuffle=False, **kw)
    te = DataLoader(ImageList(split["root"], split["test"], eval_tf), bs, shuffle=False, **kw)
    return tr, va, te


# --------------------------------------------------------------------------- #
#  models (all built from scratch, weights=None)
# --------------------------------------------------------------------------- #
def build_model(name, nc):
    if name == "SMG_Lite":
        return smg_lite(num_classes=nc, width=1.0)
    if name == "MobileNetV3-Large":
        m = models.mobilenet_v3_large(weights=None)
        m.classifier[3] = nn.Linear(m.classifier[3].in_features, nc)
        return m
    if name == "MobileNetV2":
        m = models.mobilenet_v2(weights=None)
        m.classifier[1] = nn.Linear(m.classifier[1].in_features, nc)
        return m
    if name == "VGG16":
        return models.vgg16(weights=None, num_classes=nc)
    if name == "ResNet50":
        return models.resnet50(weights=None, num_classes=nc)
    if name == "InceptionV3":
        return models.inception_v3(weights=None, aux_logits=False, init_weights=True,
                                   num_classes=nc)
    if name == "DenseNet121":
        return models.densenet121(weights=None, num_classes=nc)
    if name == "Xception":
        import timm
        for arch in ("legacy_xception", "xception"):
            try:
                return timm.create_model(arch, pretrained=False, num_classes=nc)
            except Exception:
                pass
        raise RuntimeError("Xception needs `pip install timm`")
    raise ValueError(f"unknown model {name}")


# ---- ablation: switch SMG_Lite components off WITHOUT editing SMG_Lite.py ---- #
class PlainExpand(nn.Sequential):
    """Ghost OFF -> ordinary 1x1 expansion conv (cin -> mid)."""
    def __init__(self, cin, mid):
        super().__init__(nn.Conv2d(cin, mid, 1, bias=False), nn.BatchNorm2d(mid), nn.Hardswish())


class PlainDW(nn.Sequential):
    """MS-DW gate OFF -> single depthwise 3x3 (no multi-scale, no input-dependent gate)."""
    def __init__(self, ch, stride):
        super().__init__(nn.Conv2d(ch, ch, 3, stride, 1, groups=ch, bias=False),
                         nn.BatchNorm2d(ch), nn.Hardswish())


def build_ablation(nc, ghost=True, msdw=True, cea=True):
    m = smg_lite(num_classes=nc, width=1.0)
    for blk in m.blocks:
        if not ghost and isinstance(blk.expand, GhostExpand):
            conv = blk.expand.primary[0]
            blk.expand = PlainExpand(conv.in_channels, conv.out_channels * 2)
        if not msdw:
            conv = blk.ms.dw3[0]
            blk.ms = PlainDW(conv.in_channels, conv.stride)
        if not cea and isinstance(blk.attn, CoordECA):
            blk.attn = nn.Identity()
    m._init_weights()
    return m


class WarmupEMA(EMA):
    """train.py's EMA with the usual decay warm-up: decay_t = min(decay, (1+t)/(10+t)).
    With a fixed 0.999 a short run (small dataset) keeps a large share of the random
    initial weights inside the EMA; the warm-up removes that."""

    def __init__(self, model, decay=0.999):
        super().__init__(model, decay)
        self.max_decay, self.n = decay, 0

    def update(self, model):
        self.n += 1
        self.decay = min(self.max_decay, (1 + self.n) / (10 + self.n))
        super().update(model)


def _vtag(g, m, c):
    return f"G{int(g)}_M{int(m)}_C{int(c)}"


# --------------------------------------------------------------------------- #
#  training / evaluation
# --------------------------------------------------------------------------- #
@torch.no_grad()
def _predict(model, loader, tta=False):
    model.eval()
    probs, ys = [], []
    for i, (x, y) in enumerate(loader):
        if CFG["max_batches"] and i >= CFG["max_batches"]:
            break
        x = x.to(DEVICE, non_blocking=True)
        with torch.autocast(device_type=DEVICE, enabled=(DEVICE == "cuda")):
            p = F.softmax(model(x).float(), 1)
            if tta:
                p = 0.5 * (p + F.softmax(model(torch.flip(x, dims=[3])).float(), 1))
        probs.append(p.float().cpu()); ys.append(y)
    return torch.cat(probs).numpy().astype(np.float64), torch.cat(ys).numpy()


def _loss_acc(probs, y):
    loss = float(-np.log(np.clip(probs[np.arange(len(y)), y], 1e-12, 1)).mean())
    return loss, float((probs.argmax(1) == y).mean() * 100)


def compute_metrics(probs, y, nc):
    from sklearn.metrics import (accuracy_score, precision_recall_fscore_support,
                                 roc_auc_score, cohen_kappa_score)
    probs = probs / probs.sum(1, keepdims=True)
    pred = probs.argmax(1)
    p, r, f1, _ = precision_recall_fscore_support(y, pred, average="macro", zero_division=0)
    try:
        auc = roc_auc_score(y, probs[:, 1]) if nc == 2 else \
            roc_auc_score(y, probs, multi_class="ovr", average="macro", labels=list(range(nc)))
    except ValueError:
        auc = float("nan")
    loss, _ = _loss_acc(probs, y)
    return dict(loss=loss, accuracy=accuracy_score(y, pred) * 100, precision=p * 100,
                recall=r * 100, f1=f1 * 100, auc=auc * 100,
                kappa=cohen_kappa_score(y, pred) * 100)


def model_cost(model_fn, nc):
    m = model_fn(nc).cpu()
    params = sum(p.numel() for p in m.parameters())
    trainable = sum(p.numel() for p in m.parameters() if p.requires_grad)
    size_mb = sum(p.numel() * p.element_size() for p in m.parameters()) / 1024 ** 2
    gflops = 2 * macs_of(m, CFG["img_size"]) / 1e9
    return dict(params=params, trainable=trainable, size_mb=size_mb, gflops=gflops)


def _train(run_dir, label, model_fn, ds, epochs, aug=True, lr=None, init_from=None):
    mfile = os.path.join(run_dir, "metrics.json")
    if os.path.exists(mfile) and not CFG["force"]:
        with open(mfile) as f:
            res = json.load(f)
        print(f"[skip] {label}: already done -> " + _fmt_row(res))
        return res
    os.makedirs(run_dir, exist_ok=True)
    seed_all(CFG["seed"])
    split = get_split(ds)
    nc = len(split["classes"])
    cost = model_cost(model_fn, nc)          # size / params / GFLOPs (before training: less RAM)
    tr_ld, va_ld, te_ld = _loaders(split, aug)
    model = model_fn(nc).to(DEVICE)
    if init_from:
        if os.path.exists(init_from):
            sd = torch.load(init_from, map_location=DEVICE)["model"]
            own = model.state_dict()
            keep = {k: v for k, v in sd.items() if k in own and own[k].shape == v.shape}
            model.load_state_dict(keep, strict=False)
            print(f"  initialised from {init_from} ({len(keep)}/{len(own)} tensors; new classifier)")
        else:
            print(f"  (no checkpoint at {init_from} -> training from scratch)")

    lr = lr or CFG["lr"]
    crit = nn.CrossEntropyLoss(label_smoothing=CFG["smooth"] if aug else 0.0)
    opt = torch.optim.AdamW(model.parameters(), lr=lr, weight_decay=CFG["wd"])
    steps = len(tr_ld) if not CFG["max_batches"] else min(len(tr_ld), CFG["max_batches"])
    warm, total = CFG["warmup"] * steps, epochs * steps

    def lr_at(step):  # linear warm-up then cosine decay (train.py)
        if step < warm:
            return step / max(warm, 1)
        return 0.5 * (1 + math.cos(math.pi * (step - warm) / max(total - warm, 1)))

    sched = torch.optim.lr_scheduler.LambdaLR(opt, lr_at)
    scaler = torch.amp.GradScaler(enabled=(DEVICE == "cuda"))
    ema = WarmupEMA(model, decay=CFG["ema_decay"])
    hist = {k: [] for k in ("train_loss", "train_acc", "val_loss", "val_acc")}
    best, best_loss, start = -1.0, 1e9, 0

    last = os.path.join(run_dir, "last.pt")
    if os.path.exists(last) and not CFG["force"]:          # resume an interrupted cell
        ck = torch.load(last, map_location=DEVICE, weights_only=False)
        model.load_state_dict(ck["model"]); ema.shadow.load_state_dict(ck["ema"])
        opt.load_state_dict(ck["opt"]); sched.load_state_dict(ck["sched"])
        scaler.load_state_dict(ck["scaler"]); ema.n = ck.get("ema_n", 0)
        hist, best, best_loss, start = ck["hist"], ck["best"], ck["best_loss"], ck["epoch"]
        print(f"  resumed at epoch {start}/{epochs}")

    n_par = sum(p.numel() for p in model.parameters())
    print(f"{label} | device={DEVICE} | classes={nc} | params={n_par/1e6:.3f}M | "
          f"epochs={epochs} | lr={lr} | aug={aug}")
    mix_until = int(round(epochs * (1 - CFG["mixup_off_frac"])))
    for ep in range(start, epochs):
        model.train()
        t0, run_loss, run_corr, seen = time.time(), 0.0, 0.0, 0
        use_mix = aug and CFG["mixup"] > 0 and ep < mix_until
        for i, (x, y) in enumerate(tr_ld):
            if CFG["max_batches"] and i >= CFG["max_batches"]:
                break
            x, y = x.to(DEVICE, non_blocking=True), y.to(DEVICE, non_blocking=True)
            if use_mix:
                x, ya, yb, lam = mix_batch(x, y, alpha=CFG["mixup"])
            else:
                ya, yb, lam = y, y, 1.0
            opt.zero_grad(set_to_none=True)
            with torch.autocast(device_type=DEVICE, enabled=(DEVICE == "cuda")):
                out = model(x)
                loss = lam * crit(out, ya) + (1 - lam) * crit(out, yb)
            scaler.scale(loss).backward()
            scaler.step(opt); scaler.update(); sched.step(); ema.update(model)
            with torch.no_grad():                           # plain CE for the curves
                o = out.float()
                plain = lam * F.cross_entropy(o, ya) + (1 - lam) * F.cross_entropy(o, yb)
                pr = o.argmax(1)
                run_corr += float(lam * (pr == ya).sum() + (1 - lam) * (pr == yb).sum())
            run_loss += float(plain) * y.size(0); seen += y.size(0)

        pv, yv = _predict(model, va_ld)
        pe, ye = _predict(ema.shadow, va_ld)
        (vl, va), (el, ea) = _loss_acc(pv, yv), _loss_acc(pe, ye)
        use_ema = (ea, -el) >= (va, -vl)
        v_loss, v_acc = (el, ea) if use_ema else (vl, va)
        hist["train_loss"].append(run_loss / max(seen, 1))
        hist["train_acc"].append(100 * run_corr / max(seen, 1))
        hist["val_loss"].append(v_loss); hist["val_acc"].append(v_acc)
        if (v_acc, -v_loss) > (best, -best_loss):
            best, best_loss = v_acc, v_loss
            torch.save({"model": (ema.shadow if use_ema else model).state_dict(),
                        "epoch": ep + 1, "val_acc": v_acc}, os.path.join(run_dir, "best.pt"))
            _backup(os.path.join(run_dir, "best.pt"))
        torch.save({"model": model.state_dict(), "ema": ema.shadow.state_dict(),
                    "opt": opt.state_dict(), "sched": sched.state_dict(),
                    "scaler": scaler.state_dict(), "hist": hist, "best": best,
                    "best_loss": best_loss, "epoch": ep + 1, "ema_n": ema.n}, last)
        _backup(last)
        print(f"  ep {ep+1:3d}/{epochs}  train_loss {hist['train_loss'][-1]:.4f}  "
              f"train_acc {hist['train_acc'][-1]:6.2f}  val_loss {v_loss:.4f}  "
              f"val_acc {v_acc:6.2f}{' (ema)' if use_ema else ''}  best {best:6.2f}  "
              f"{time.time()-t0:5.1f}s")

    # ---- test the best-on-validation checkpoint once ----
    ck = torch.load(os.path.join(run_dir, "best.pt"), map_location=DEVICE)
    model.load_state_dict(ck["model"])
    probs, y = _predict(model, te_ld, tta=CFG["tta"])
    res = compute_metrics(probs, y, nc)
    res.update(cost)
    res.update(label=label, dataset=ds, epochs=epochs, best_epoch=ck["epoch"],
               best_val_acc=ck["val_acc"], aug=aug, n_test=int(len(y)), classes=split["classes"])
    np.savez_compressed(os.path.join(run_dir, "preds.npz"), probs=probs.astype(np.float32), y=y)
    with open(os.path.join(run_dir, "history.json"), "w") as f:
        json.dump(hist, f)
    with open(mfile, "w") as f:
        json.dump(res, f, indent=1)
    if os.path.exists(last):
        os.remove(last)
    _unbackup(last)
    _backup(run_dir)                      # metrics.json, history.json, preds.npz, best.pt
    print("TEST -> " + _fmt_row(res))
    del model, ema, opt, sched, scaler, ck, tr_ld, va_ld, te_ld   # free RAM/GPU for the next cell
    import gc; gc.collect()
    if DEVICE == "cuda":
        torch.cuda.empty_cache()
    return res


def _fmt_row(r):
    return (f"loss {r['loss']:.4f} | acc {r['accuracy']:.2f}% | prec {r['precision']:.2f}% | "
            f"rec {r['recall']:.2f}% | f1 {r['f1']:.2f}% | auc {r['auc']:.2f}% | "
            f"kappa {r['kappa']:.2f}% | params {r['params']:,} | {r['size_mb']:.2f} MB | "
            f"{r['gflops']:.3f} GFLOPs")


def run(dataset, model, aug=True):
    """One cell = one model trained & tested on one dataset."""
    tag = model if aug else f"{model}__noaug"
    init = None
    src = CFG.get("transfer_from")
    if src and src != dataset:
        init = os.path.join(CFG["out_dir"], "runs", src, model, "best.pt")
    lr = CFG["lr_override"].get(model)
    return _train(os.path.join(CFG["out_dir"], "runs", dataset, tag),
                  f"[{dataset}] {tag}", lambda nc: build_model(model, nc), dataset,
                  CFG["datasets"][dataset]["epochs"], aug=aug, lr=lr, init_from=init)


def run_ablation(dataset, ghost, msdw, cea, epochs=None):
    """One cell = one ablation variant (own epoch budget, trained from scratch)."""
    tag = _vtag(ghost, msdw, cea)
    return _train(os.path.join(CFG["out_dir"], "ablation", dataset, tag),
                  f"[ablation {dataset}] Ghost={ghost} MS-DW gate={msdw} CEA={cea}",
                  lambda nc: build_ablation(nc, ghost, msdw, cea), dataset,
                  epochs or CFG["ablation_epochs"])


# --------------------------------------------------------------------------- #
#  combine: tables
# --------------------------------------------------------------------------- #
def _load(run_dir):
    f = os.path.join(run_dir, "metrics.json")
    if not os.path.exists(f):
        return None
    with open(f) as fh:
        return json.load(fh)


def _cells(r):
    if r is None:
        return [""] * 10
    pct = lambda v: "" if v is None or (isinstance(v, float) and math.isnan(v)) else f"{v:.2f} %"
    return [f"{r['size_mb']:.2f}", f"{r['trainable']:,}", f"{r['loss']:.4f}",
            pct(r["accuracy"]), pct(r["precision"]), pct(r["recall"]), pct(r["f1"]),
            pct(r["auc"]), pct(r["kappa"]), f"{r['gflops']:.3f}"]


def _best_mask(results, keys=("loss", "accuracy", "precision", "recall", "f1", "auc", "kappa")):
    """(row, metric) pairs to print in bold: lowest loss, highest of everything else."""
    bold = set()
    for k in keys:
        vals = [(i, r[k]) for i, r in enumerate(results)
                if r is not None and not (isinstance(r[k], float) and math.isnan(r[k]))]
        if len(vals) < 2 or len({round(v, 4) for _, v in vals}) == 1:
            continue
        tgt = min(v for _, v in vals) if k == "loss" else max(v for _, v in vals)
        bold |= {(i, k) for i, v in vals if round(v, 4) == round(tgt, 4)}
    return bold


def results_table(dataset):
    import pandas as pd
    res = [_load(os.path.join(CFG["out_dir"], "runs", dataset, m)) for m in MODEL_NAMES]
    rows = [[DISPLAY[m]] + _cells(r) for m, r in zip(MODEL_NAMES, res)]
    return pd.DataFrame(rows, columns=TABLE_COLS), res


def ablation_table(dataset):
    import pandas as pd
    res = [_load(os.path.join(CFG["out_dir"], "ablation", dataset, _vtag(*v)))
           for v in ABLATION_VARIANTS]
    rows = []
    for (g, m, c), r in zip(ABLATION_VARIANTS, res):
        cells = _cells(r)
        rows.append(["✓" if g else "–", "✓" if m else "–",
                     "✓" if c else "–", cells[0],
                     "" if r is None else f"{r['params']:,}"] + cells[2:9])
    cols = ["Ghost", "MS-DW gate", "CEA", "Model Size", "Total parameters", "Loss", "Accuracy",
            "Precision", "Recall", "F1-score", "AUC-Score", "Kappa Score"]
    return pd.DataFrame(rows, columns=cols), res


_METRIC_OF_COL = {"Loss": "loss", "Accuracy": "accuracy", "Precision": "precision",
                  "Recall": "recall", "F1-score": "f1", "AUC- Score": "auc", "AUC-Score": "auc",
                  "Kappa Score": "kappa"}


def table_png(df, results, path, title=None, check_cols=0):
    """Booktabs-style table image (same look as the ablation table in the docx)."""
    import matplotlib.pyplot as plt
    bold = _best_mask(results)
    ncol, nrow = len(df.columns), len(df)
    widths = [max(len(str(c)), *(len(str(v)) for v in df[c])) + 3 for c in df.columns]
    xs = np.cumsum([0] + widths[:-1]) / sum(widths)
    fig_w = max(10, sum(widths) * 0.115)
    fig = plt.figure(figsize=(fig_w, 0.38 * (nrow + 1.5) + (0.4 if title else 0)), dpi=200)
    ax = fig.add_axes([0.01, 0, 0.98, 0.9 if title else 1]); ax.axis("off")
    ax.set_xlim(0, 1); ax.set_ylim(nrow + 1.25, 0)
    if title:
        fig.suptitle(title, fontsize=12, family="DejaVu Serif")
    ax.plot([0, 1], [0.15, 0.15], "k", lw=1.2); ax.plot([0, 1], [1.05, 1.05], "k", lw=0.8)
    ax.plot([0, 1], [nrow + 1.15, nrow + 1.15], "k", lw=1.2)
    for x, c in zip(xs, df.columns):
        ax.text(x, 0.6, c, fontsize=10, va="center", family="DejaVu Serif")
    for i in range(nrow):
        for j, (x, c) in enumerate(zip(xs, df.columns)):
            v = str(df.iloc[i, j])
            is_check = j < check_cols
            b = (i, _METRIC_OF_COL.get(c)) in bold
            ax.text(x + (0.01 if is_check else 0), i + 1.55, v, fontsize=11 if is_check else 10,
                    va="center", family="DejaVu Sans" if is_check else "DejaVu Serif",
                    fontweight="bold" if b else "normal")
    fig.savefig(path, bbox_inches="tight", pad_inches=0.12, facecolor="white")
    plt.close(fig)
    return path


# --------------------------------------------------------------------------- #
#  combine: figures (same style as the images in result_word.docx)
# --------------------------------------------------------------------------- #
def _title(ds):
    return CFG["datasets"][ds].get("title", ds)


def plot_metric_lines(dataset, path):
    import matplotlib.pyplot as plt
    res = [(DISPLAY[m], _load(os.path.join(CFG["out_dir"], "runs", dataset, m)))
           for m in MODEL_NAMES]
    res = [(n, r) for n, r in res if r is not None]
    if not res:
        return None
    fig, ax = plt.subplots(figsize=(11, 5), dpi=150)
    x = np.arange(len(res))
    for key, lab in (("accuracy", "Accuracy (%)"), ("precision", "Precision (%)"),
                     ("recall", "Recall (%)"), ("f1", "F1-score (%)")):
        ax.plot(x, [r[key] for _, r in res], marker="o", label=lab)
    ax.set_xticks(x); ax.set_xticklabels([n for n, _ in res], rotation=45)
    ax.set_ylabel("Performance Metrics (%)", fontsize=11)
    ax.set_title(f"Model Performance Metrics Comparison ({_title(dataset)})", fontsize=13)
    ax.legend(loc="upper left")
    fig.tight_layout(); fig.savefig(path, facecolor="white"); plt.close(fig)
    return path


def plot_curves(dataset, path, baseline="MobileNetV2", ours="SMG_Lite"):
    import matplotlib.pyplot as plt
    pal = ["#4C72B0", "#DD8452", "#55A868", "#C44E52", "#8172B3", "#937860", "#DA8BC3", "#8C8C8C"]
    # same 8 lines as the template: no-aug runs first, then the "(Aug)" runs
    runs = [(baseline, f"{baseline}__noaug"), ("Proposed Model (SMG_Lite)", f"{ours}__noaug"),
            (baseline, baseline), ("Proposed Model (SMG_Lite)", ours)]
    hs = []
    for k, (name, tag) in enumerate(runs):
        f = os.path.join(CFG["out_dir"], "runs", dataset, tag, "history.json")
        if os.path.exists(f):
            with open(f) as fh:
                hs.append((name, k >= 2, json.load(fh), k))
    if not hs:
        return None
    with plt.style.context("seaborn-v0_8-whitegrid" if "seaborn-v0_8-whitegrid"
                           in plt.style.available else "default"):
        fig, axes = plt.subplots(2, 1, figsize=(9, 9), dpi=150)
        for ax, key, ttl, ylab in ((axes[0], "acc", "Accuracy", "Accuracy (%)"),
                                   (axes[1], "loss", "Loss", "Loss")):
            for name, is_aug, h, k in hs:
                sfx = " (Aug)" if is_aug else ""
                ep = np.arange(1, len(h[f"train_{key}"]) + 1)
                ax.plot(ep, h[f"train_{key}"], "-o", ms=4, color=pal[2 * k],
                        label=f"{name} Training{sfx}")
                ax.plot(ep, h[f"val_{key}"], "--x", ms=5, color=pal[2 * k + 1],
                        label=f"{name} Validation{sfx}")
            ax.set_title(ttl); ax.set_xlabel("Epochs"); ax.set_ylabel(ylab)
            ax.xaxis.set_major_locator(plt.MaxNLocator(integer=True))
            ax.legend(loc="lower right" if key == "acc" else "upper right", fontsize=8,
                      frameon=True)
        fig.suptitle(_title(dataset), fontsize=12)
        fig.tight_layout(); fig.savefig(path, facecolor="white"); plt.close(fig)
    return path


def _preds(dataset, model="SMG_Lite"):
    d = os.path.join(CFG["out_dir"], "runs", dataset, model)
    if not os.path.exists(os.path.join(d, "preds.npz")):
        return None
    z = np.load(os.path.join(d, "preds.npz"))
    return z["probs"].astype(np.float64), z["y"], _load(d)["classes"]


def plot_confusion(dataset, path, model="SMG_Lite"):
    import matplotlib.pyplot as plt
    from sklearn.metrics import confusion_matrix
    pr = _preds(dataset, model)
    if pr is None:
        return None
    probs, y, classes = pr
    nc = len(classes)
    cm = confusion_matrix(y, probs.argmax(1), labels=list(range(nc)))
    acc = cm.diagonal() / np.maximum(cm.sum(1), 1) * 100
    s = max(7, 0.38 * nc + 4)
    fig, ax = plt.subplots(figsize=(s * 1.15, s), dpi=150)
    im = ax.imshow(cm, cmap="Blues")
    fs = 9 if nc <= 10 else (7 if nc <= 20 else 5.5)
    thr = cm.max() / 2
    for i in range(nc):
        for j in range(nc):
            ax.text(j, i, cm[i, j], ha="center", va="center", fontsize=fs,
                    color="white" if cm[i, j] > thr else "black")
    ax.set_xticks(range(nc)); ax.set_yticks(range(nc))
    ax.set_xticklabels([f"{i}, Acc: {a:.2f}%" for i, a in enumerate(acc)], rotation=45,
                       ha="right", fontsize=fs + 1)
    ax.set_yticklabels(range(nc), fontsize=fs + 1)
    ax.set_xlabel("Predicted label (Accuracy)"); ax.set_ylabel("True label")
    ax.set_title(f"Confusion Matrix ({_title(dataset)})")
    fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
    fig.tight_layout(); fig.savefig(path, facecolor="white"); plt.close(fig)
    with open(os.path.splitext(path)[0] + "_class_index.csv", "w") as f:
        f.write("index,class,test_images,class_accuracy\n")
        for i, c in enumerate(classes):
            f.write(f"{i},{c},{cm[i].sum()},{acc[i]:.2f}\n")
    return path


def plot_roc(dataset, path, model="SMG_Lite"):
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve, auc
    pr = _preds(dataset, model)
    if pr is None:
        return None
    probs, y, classes = pr
    fig, ax = plt.subplots(figsize=(10, 9), dpi=150)
    cmap = plt.get_cmap("tab20" if len(classes) > 10 else "tab10")
    for i, c in enumerate(classes):
        if (y == i).sum() == 0:
            continue
        fpr, tpr, _ = roc_curve(y == i, probs[:, i])
        ax.plot(fpr, tpr, lw=1.5, ls="--", color=cmap(i % cmap.N),
                label=f"{i}: {c} (AUC = {auc(fpr, tpr):.4f})")
    ax.plot([0, 1], [0, 1], "k--", lw=1.5)
    ax.set_xlim(0, 1); ax.set_ylim(0, 1.05); ax.grid(True)
    ax.set_xlabel("False Positive Rate"); ax.set_ylabel("True Positive Rate")
    ax.set_title("ROC Curves for All Classes")
    fig.suptitle(f"{_title(dataset)} Dataset - ROC Curve", fontsize=16)
    if len(classes) <= 12:
        ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout(); fig.savefig(path, facecolor="white"); plt.close(fig)
    return path


# --------------------------------------------------------------------------- #
#  combine: fill templates/result_word.docx
# --------------------------------------------------------------------------- #
def _placeholder(path, text):
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(8, 1.2), dpi=120)
    fig.text(0.5, 0.5, text, ha="center", va="center", fontsize=12, color="#888")
    fig.savefig(path, facecolor="white"); plt.close(fig)
    return path


def _fit_columns(tbl, df):
    """Re-distribute the template's column widths by content so numbers do not wrap."""
    from docx.shared import Emu
    grid = tbl._tbl.tblGrid.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol")
    total = sum(int(g.get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w") or 0) for g in grid)
    if not total or len(grid) != len(df.columns):
        return
    need = []
    for j, c in enumerate(df.columns):
        longest_word = max(len(w) for w in str(c).split()) * 1.25   # header is 9 pt bold-ish
        need.append(max([longest_word] + [len(str(v)) for v in df.iloc[:, j]]) + 1.5)
    tbl.autofit = False
    for j, g in enumerate(grid):
        w = int(total * need[j] / sum(need))
        g.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", str(w))
        for row in tbl.rows:
            row.cells[j].width = Emu(w * 635)        # twips -> EMU


def fill_docx(template, out_path, table_datasets, figures, ablation_datasets):
    """
    table_datasets : dataset names in the order of the 3 tables of the template
    figures        : list of 5 lists of (png, caption) -- one list per template image
    """
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(template)

    # ---- tables ----
    for tbl, ds in zip(doc.tables, table_datasets):
        df, res = results_table(ds)
        bold = _best_mask(res)
        hdr_runs = [r for p in tbl.rows[1].cells[0].paragraphs for r in p.runs]
        size = hdr_runs[0].font.size if hdr_runs else None
        for i, row in enumerate(tbl.rows[1:]):
            if i >= len(df):
                break
            for j, cell in enumerate(row.cells):
                val = str(df.iloc[i, j])
                para = cell.paragraphs[0]
                if j == 0:
                    if para.runs:
                        para.runs[0].text = val
                        for r in para.runs[1:]:
                            r.text = ""
                    else:
                        para.add_run(val)
                    continue
                for r in para.runs:
                    r.text = ""
                run = para.add_run(val)
                run.font.size = Pt(8.5)
                run.bold = (i, _METRIC_OF_COL.get(df.columns[j])) in bold
        for cell in tbl.rows[0].cells:                 # header: 9 pt so words do not split
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
        _fit_columns(tbl, df)

    # ---- captions ----
    names = ", ".join(_title(d) for d in table_datasets)
    for p in doc.paragraphs:
        t = p.text.strip()
        new = None
        if t.startswith("Three Performance matrix"):
            new = f"Three performance matrices (training/validation curves): {names}"
        elif t.startswith("Three separate confusion"):
            new = f"Three separate confusion matrices of SMG_Lite: {names}"
        elif t.startswith("Three separate ROC"):
            new = f"Three separate ROCs of SMG_Lite: {names}"
        elif t.startswith("Alation studies") or t.startswith("Ablation studies"):
            def _ep(d):
                rs = [r for r in ablation_table(d)[1] if r]
                return rs[0]["epochs"] if rs else CFG["ablation_epochs"]
            new = "Ablation studies on proposed model (" + \
                  ", ".join(f"{_title(d)}, {_ep(d)} epochs" for d in ablation_datasets) + ")"
        if new:
            for r in p.runs[1:]:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)

    # ---- images: replace each template picture by our figure(s) ----
    pic_paras = [p for p in doc.paragraphs if "graphicData" in p._p.xml]
    for para, figs in zip(pic_paras, figures):
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        anchor = para
        for k, (png, cap) in enumerate(figs):
            tgt = para if k == 0 else anchor
            if k > 0:
                new_p = OxmlElement("w:p"); anchor._p.addnext(new_p)
                tgt = Paragraph(new_p, para._parent)
            tgt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tgt.add_run().add_picture(png, width=Inches(6.5))
            anchor = tgt
            if cap:
                new_p = OxmlElement("w:p"); anchor._p.addnext(new_p)
                cp = Paragraph(new_p, para._parent)
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cp.add_run(cap); run.italic = True; run.font.size = Pt(9)
                anchor = cp
    doc.save(out_path)
    return out_path


def make_report(table_datasets, ablation_datasets, template="templates/result_word.docx"):
    """Build every table, figure and the filled Word file from the saved results."""
    from IPython.display import display, Image as IImage, Markdown
    tdir, fdir = _p("tables", "x"), _p("figures", "x")
    tdir, fdir = os.path.dirname(tdir), os.path.dirname(fdir)
    letters = "abcdefghij"
    slots = [[], [], [], [], []]
    for k, ds in enumerate(table_datasets):
        df, res = results_table(ds)
        df.to_csv(os.path.join(tdir, f"table_{ds}.csv"), index=False)
        display(Markdown(f"### Performance comparison of various Models on the {_title(ds)} dataset"))
        display(df)
        png = table_png(df, res, os.path.join(tdir, f"table_{ds}.png"))
        cap = f"({letters[k]}) {_title(ds)}"
        for slot, fn, name in ((0, plot_metric_lines, "metrics_lines"), (1, plot_curves, "curves"),
                               (2, plot_confusion, "confusion"), (3, plot_roc, "roc")):
            f = fn(ds, os.path.join(fdir, f"{name}_{ds}.png"))
            slots[slot].append((f or _placeholder(os.path.join(fdir, f"{name}_{ds}_missing.png"),
                                                  f"{_title(ds)}: run the experiments first"), cap))
    for k, ds in enumerate(ablation_datasets):
        df, res = ablation_table(ds)
        df.to_csv(os.path.join(tdir, f"ablation_{ds}.csv"), index=False)
        display(Markdown(f"### Ablation studies on proposed model ({_title(ds)})"))
        display(df)
        png = table_png(df, res, os.path.join(tdir, f"ablation_{ds}.png"), check_cols=3)
        slots[4].append((png, f"({letters[k]}) {_title(ds)}" if len(ablation_datasets) > 1 else ""))
    out = _p("result_word_filled.docx")
    fill_docx(template, out, table_datasets, slots, ablation_datasets)
    for d in (tdir, fdir, out):
        _backup(d)
    print(f"\nWord report written -> {out}")
    print(f"Tables (csv/png)    -> {tdir}\nFigures (png)       -> {fdir}")
    if CFG.get("backup_dir"):
        print(f"Everything is also saved in {CFG['backup_dir']}")
    return out
